#!/usr/bin/env python3
"""生成葡萄酒论文PPT"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ── 色彩方案：酒红色系 ──
C_WINE    = RGBColor(0x72, 0x24, 0x2F)   # 酒红主色
C_WINE2   = RGBColor(0x5C, 0x1A, 0x24)   # 深酒红
C_GOLD    = RGBColor(0xC8, 0xA2, 0x5C)   # 金色点缀
C_CREAM   = RGBColor(0xF5, 0xEE, 0xE0)   # 奶油底
C_DARK    = RGBColor(0x2C, 0x18, 0x1C)   # 深色字
C_GRAY    = RGBColor(0x88, 0x77, 0x72)   # 灰色辅助文字
C_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
C_LIGHT   = RGBColor(0xFA, 0xF3, 0xE8)   # 浅暖色
C_BG_DARK = RGBColor(0x3E, 0x22, 0x28)   # 深色页背景

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)

W = prs.slide_width
H = prs.slide_height

def add_bg(slide, color):
    """设置幻灯片纯色背景"""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_shape_bg(slide, color, left=0, top=0, width=None, height=None):
    """添加色块形状"""
    w = width or W
    h = height or H
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_textbox(slide, left, top, width, height, text, font_size=18,
                color=C_DARK, bold=False, alignment=PP_ALIGN.LEFT,
                font_name='Microsoft YaHei', line_spacing=1.3):
    """添加文本框"""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    p.space_after = Pt(0)
    p.space_before = Pt(0)
    if line_spacing:
        p.line_spacing = Pt(font_size * line_spacing)
    return txBox

def add_para(text_frame, text, font_size=18, color=C_DARK, bold=False,
             alignment=PP_ALIGN.LEFT, font_name='Microsoft YaHei',
             space_before=0, space_after=0, line_spacing=None):
    """在已有text_frame上添加段落"""
    p = text_frame.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    if line_spacing:
        p.line_spacing = Pt(font_size * line_spacing)
    return p

def add_accent_bar(slide, left, top, width=Inches(0.08), height=Inches(1.5), color=C_GOLD):
    """添加装饰条"""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_bottom_line(slide, color=C_GOLD, opacity=True):
    """底部装饰线"""
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0.8), H - Inches(0.6),
        W - Inches(1.6), Pt(2)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape

def add_page_number(slide, num, total=10):
    add_textbox(slide, W - Inches(1.2), H - Inches(0.5),
                Inches(1.0), Inches(0.4),
                f"{num}/{total}", font_size=10, color=C_GRAY,
                alignment=PP_ALIGN.RIGHT, font_name='Arial')

# ════════════════════════════════════════════
# Slide 1: 封面
# ════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
add_bg(slide, C_BG_DARK)

# 装饰 top bar
add_shape_bg(slide, C_GOLD, top=0, height=Pt(6))

# 大标题
add_textbox(slide, Inches(1.5), Inches(1.8), Inches(10), Inches(1.5),
            "液体记忆", font_size=54, color=C_GOLD, bold=True,
            font_name='Microsoft YaHei')

# 副标题
add_textbox(slide, Inches(1.5), Inches(3.4), Inches(10), Inches(1.0),
            "LUQUID MEMORY", font_size=24, color=C_GOLD,
            alignment=PP_ALIGN.LEFT, font_name='Arial')

# 中文副标题
add_textbox(slide, Inches(1.5), Inches(4.2), Inches(10), Inches(0.8),
            "论葡萄酒作为历史、风土与情感的复合载体",
            font_size=22, color=C_CREAM, font_name='Microsoft YaHei')

# 分隔线
shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(5.3), Inches(3), Pt(2))
shape.fill.solid()
shape.fill.fore_color.rgb = C_GOLD
shape.line.fill.background()

# 作者信息
add_textbox(slide, Inches(1.5), Inches(5.6), Inches(6), Inches(0.5),
            "跨学科研究 · 文化哲学 · 2026",
            font_size=14, color=C_GRAY, font_name='Microsoft YaHei')

add_textbox(slide, Inches(1.5), Inches(6.1), Inches(6), Inches(0.4),
            "基于同名学术论文改编",
            font_size=11, color=C_GRAY, font_name='Microsoft YaHei')

# 右下角装饰 - 小方块
for i in range(3):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
        W - Inches(1.5) + i * Inches(0.18), H - Inches(0.35),
        Inches(0.12), Inches(0.12))
    shape.fill.solid()
    shape.fill.fore_color.rgb = C_GOLD if i == 0 else C_GRAY
    shape.line.fill.background()

add_page_number(slide, 1)

# ════════════════════════════════════════════
# Slide 2: 目录
# ════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, C_CREAM)
add_bottom_line(slide)

add_textbox(slide, Inches(1.5), Inches(0.6), Inches(4), Inches(0.5),
            "CONTENTS", font_size=14, color=C_GOLD, bold=True,
            font_name='Arial', alignment=PP_ALIGN.LEFT)
add_textbox(slide, Inches(1.5), Inches(1.0), Inches(6), Inches(0.7),
            "目录", font_size=36, color=C_WINE, bold=True,
            font_name='Microsoft YaHei')

# 目录项
toc_items = [
    ("01", "引言：被忽视的叙事维度"),
    ("02", "自然的档案：葡萄酒作为风土的精确记录"),
    ("03", "文化的化石：传统与技艺的非物质遗产"),
    ("04", "情感的载体：从个人记忆到集体无意识"),
    ("05", "结论：一种新的阅读方式"),
]

for i, (num, title) in enumerate(toc_items):
    y = Inches(2.2) + i * Inches(0.9)
    # 数字
    add_textbox(slide, Inches(1.5), y, Inches(0.8), Inches(0.6),
                num, font_size=28, color=C_GOLD, bold=True,
                font_name='Arial')
    # 装饰竖线
    add_accent_bar(slide, Inches(2.3), y + Pt(4), Pt(3), Inches(0.45), C_GOLD)
    # 标题
    add_textbox(slide, Inches(2.6), y + Pt(2), Inches(9), Inches(0.5),
                title, font_size=20, color=C_DARK,
                font_name='Microsoft YaHei')

add_page_number(slide, 2)

# ════════════════════════════════════════════
# Slide 3: 引言
# ════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, C_CREAM)
add_bottom_line(slide)

add_textbox(slide, Inches(1.5), Inches(0.5), Inches(4), Inches(0.5),
            "01 · INTRODUCTION", font_size=12, color=C_GOLD, bold=True,
            font_name='Arial')
add_textbox(slide, Inches(1.5), Inches(0.9), Inches(6), Inches(0.7),
            "引言：被忽视的叙事维度", font_size=30, color=C_WINE, bold=True,
            font_name='Microsoft YaHei')

# 左侧 - 主要金句
add_accent_bar(slide, Inches(1.5), Inches(2.0), Pt(4), Inches(2.2), C_WINE)
add_textbox(slide, Inches(1.9), Inches(2.0), Inches(4.5), Inches(2.2),
            "当我们谈论一瓶1982年的拉菲时，谈论的远不止是赤霞珠的单宁结构和橡木桶带来的香草气息。我们实际上是在尝试触碰一个特定的时空节点。",
            font_size=18, color=C_DARK, line_spacing=1.6,
            font_name='Microsoft YaHei')

# 右侧 - 核心论点
txBox = add_textbox(slide, Inches(7.5), Inches(2.0), Inches(5), Inches(4.5),
                    "", font_size=16, color=C_DARK)
tf = txBox.text_frame
tf.word_wrap = True

add_para(tf, "核心论点", font_size=13, color=C_GOLD, bold=True,
         font_name='Arial', space_after=8)
add_para(tf, '葡萄酒的本质是一种“液体记忆”——一种能够同时记录自然环境信息、保存人类文化技艺、并激活情感体验的复合载体。',
         font_size=16, color=C_DARK, space_after=12)

add_para(tf, "研究维度", font_size=13, color=C_GOLD, bold=True,
         font_name='Arial', space_after=8)
add_para(tf, "自然的档案（风土）\n文化的化石（传统）\n情感的载体（记忆）",
         font_size=16, color=C_DARK, space_after=8, line_spacing=1.5)

add_textbox(slide, Inches(1.5), Inches(6.3), Inches(10), Inches(0.4),
            "关键词：葡萄酒 | 风土 | 液体记忆 | 文化载体 | 历史文献",
            font_size=12, color=C_GRAY, font_name='Microsoft YaHei')

add_page_number(slide, 3)

# ════════════════════════════════════════════
# Slide 4: 风土 1 - 年份的气象报告
# ════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, C_BG_DARK)

# 顶部装饰
add_shape_bg(slide, C_GOLD, top=0, height=Pt(4))

add_textbox(slide, Inches(1.5), Inches(0.5), Inches(4), Inches(0.5),
            "02 · TERROR", font_size=12, color=C_GOLD, bold=True,
            font_name='Arial')
add_textbox(slide, Inches(1.5), Inches(0.9), Inches(10), Inches(0.7),
            "自然的档案：葡萄酒作为风土的精确记录",
            font_size=30, color=C_GOLD, bold=True,
            font_name='Microsoft YaHei')

# 左列: 年份与气候
txBox = add_textbox(slide, Inches(1.5), Inches(2.0), Inches(5.5), Inches(4.8),
                    "", font_size=16, color=C_CREAM)
tf = txBox.text_frame
tf.word_wrap = True
add_para(tf, "年份：气候的气象报告", font_size=22, color=C_GOLD,
         bold=True, font_name='Microsoft YaHei', space_after=10)
add_para(tf, "凉爽多雨的年份→酸度更高、酒体更轻盈\n炎热干旱的年份→更成熟、更浓郁",
         font_size=17, color=C_CREAM, line_spacing=1.6, space_after=12)
add_para(tf, "2003年欧洲热浪期间，波尔多葡萄成熟速度空前加快，当年葡萄酒的潜在酒精度比长期平均值高出2-3个百分点。",
         font_size=15, color=C_CREAM, line_spacing=1.5, space_after=8)
add_para(tf, "通过分析葡萄酒中稳定同位素比值，科学家可以反推当年产区的气候特征。",
         font_size=15, color=C_CREAM, line_spacing=1.5)

# 右列: 土壤地质
txBox2 = add_textbox(slide, Inches(7.5), Inches(2.0), Inches(5), Inches(4.8),
                     "", font_size=16, color=C_CREAM)
tf2 = txBox2.text_frame
tf2.word_wrap = True
add_para(tf2, "土壤：可品尝的地质学", font_size=22, color=C_GOLD,
         bold=True, font_name='Microsoft YaHei', space_after=10)
add_para(tf2, "勃艮第特级园系统建立在对侏罗纪海洋沉积物分布的精细理解之上。",
         font_size=17, color=C_CREAM, line_spacing=1.6, space_after=12)

# 表格样式数据
geology_items = [
    "石灰岩  →  勃艮第 · 结构感",
    "板  岩  →  摩泽尔 · 矿物感",
    "黏  土  →  波尔多右岸 · 饱满",
    "花岗岩  →  博若莱 · 芳香",
    "火山岩  →  西西里 · 烟熏",
]
for item in geology_items:
    add_para(tf2, item, font_size=15, color=C_CREAM,
             font_name='Microsoft YaHei', line_spacing=1.5)

add_bottom_line(slide, C_GOLD)
add_page_number(slide, 4)

# ════════════════════════════════════════════
# Slide 5: 风土 2 — 数据页
# ════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, C_CREAM)
add_bottom_line(slide)

add_textbox(slide, Inches(1.5), Inches(0.5), Inches(4), Inches(0.5),
            "02 · TERROIR DATA", font_size=12, color=C_GOLD, bold=True,
            font_name='Arial')
add_textbox(slide, Inches(1.5), Inches(0.9), Inches(10), Inches(0.7),
            "风土关键词", font_size=30, color=C_WINE, bold=True,
            font_name='Microsoft YaHei')

# 四列数据块
col_data = [
    ("8000+", "年", "葡萄栽培历史"),
    ("1368", "种", "商用酿酒葡萄"),
    ("77", "国", "葡萄酒产国"),
    ("800+", "种", "芳香化合物"),
]
for i, (nb, unit, desc) in enumerate(col_data):
    x = Inches(1.3) + i * Inches(3.0)
    y = Inches(2.2)

    # 装饰条
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
        x, y, Inches(0.06), Inches(1.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = C_WINE
    shape.line.fill.background()

    # 数字
    add_textbox(slide, x + Inches(0.3), y - Inches(0.1), Inches(2.5), Inches(0.9),
                nb, font_size=44, color=C_WINE, bold=True, font_name='Arial')

    # 单位+描述
    add_textbox(slide, x + Inches(0.3), y + Inches(0.7), Inches(2.5), Inches(0.4),
                unit, font_size=16, color=C_GRAY, font_name='Microsoft YaHei')
    add_textbox(slide, x + Inches(0.3), y + Inches(1.0), Inches(2.5), Inches(0.4),
                desc, font_size=14, color=C_DARK, font_name='Microsoft YaHei')

# 底部要点
txBox = add_textbox(slide, Inches(1.5), Inches(4.2), Inches(10), Inches(2.5),
                    "", font_size=16, color=C_DARK)
tf = txBox.text_frame
tf.word_wrap = True
add_para(tf, "·   风土不应神秘化——它是可测量、可验证的环境变量在葡萄酒中的综合表达",
         font_size=17, color=C_DARK, font_name='Microsoft YaHei', space_after=6)
add_para(tf, "·   葡萄对环境极为敏感，果实化学成分精确反映生长季节的环境条件",
         font_size=17, color=C_DARK, font_name='Microsoft YaHei', space_after=6)
add_para(tf, "·   品鉴一款特定产区的葡萄酒，是一次跨越时空的地质考察",
         font_size=17, color=C_DARK, font_name='Microsoft YaHei')

add_page_number(slide, 5)

# ════════════════════════════════════════════
# Slide 6: 文化化石 开篇
# ════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, C_BG_DARK)
add_shape_bg(slide, C_GOLD, top=0, height=Pt(4))

add_textbox(slide, Inches(1.5), Inches(0.5), Inches(4), Inches(0.5),
            "03 · CULTURAL FOSSIL", font_size=12, color=C_GOLD, bold=True,
            font_name='Arial')
add_textbox(slide, Inches(1.5), Inches(0.9), Inches(10), Inches(0.7),
            "文化的化石：传统与技艺的非物质遗产",
            font_size=30, color=C_GOLD, bold=True,
            font_name='Microsoft YaHei')

txBox = add_textbox(slide, Inches(1.5), Inches(2.0), Inches(10), Inches(2.0),
                    "", font_size=16, color=C_CREAM)
tf = txBox.text_frame
tf.word_wrap = True
add_para(tf, "在许多葡萄酒产区，酿造方法本身就是活着的传统，是代代相传的非物质文化遗产。",
         font_size=20, color=C_GOLD, font_name='Microsoft YaHei',
         space_after=16, line_spacing=1.5)

# 两列对比
# 左列 - 格鲁吉亚
y = Inches(4.0)
add_accent_bar(slide, Inches(1.5), y, Pt(4), Inches(1.8), C_GOLD)
txBox = add_textbox(slide, Inches(1.9), y, Inches(5), Inches(2.0),
                    "", font_size=16, color=C_CREAM)
tf = txBox.text_frame
tf.word_wrap = True
add_para(tf, "格鲁吉亚 · Qvevri 陶罐法", font_size=20, color=C_GOLD,
         bold=True, font_name='Microsoft YaHei', space_after=8)
add_para(tf, "8000年来使用埋入地下的巨型陶罐自然发酵。品尝这样一款酒，品尝的是人类早期文明的智慧样本。",
         font_size=16, color=C_CREAM, line_spacing=1.5)

# 右列 - 香槟
add_accent_bar(slide, Inches(7.5), y, Pt(4), Inches(1.8), C_GOLD)
txBox = add_textbox(slide, Inches(7.9), y, Inches(5), Inches(2.0),
                    "", font_size=16, color=C_CREAM)
tf = txBox.text_frame
tf.word_wrap = True
add_para(tf, "香槟 · 意外演变为文化符号", font_size=20, color=C_GOLD,
         bold=True, font_name='Microsoft YaHei', space_after=8)
add_para(tf, "17世纪末瓶内二次发酵的气泡曾被视作缺陷，如今成为全球庆祝仪式的代名词。",
         font_size=16, color=C_CREAM, line_spacing=1.5)

add_bottom_line(slide, C_GOLD)
add_page_number(slide, 6)

# ════════════════════════════════════════════
# Slide 7: 传统 vs 现代
# ════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, C_CREAM)
add_bottom_line(slide)

add_textbox(slide, Inches(1.5), Inches(0.5), Inches(4), Inches(0.5),
            "03 · TRADITION VS MODERN", font_size=12, color=C_GOLD, bold=True,
            font_name='Arial')
add_textbox(slide, Inches(1.5), Inches(0.9), Inches(10), Inches(0.7),
            "现代性与传统的辩证", font_size=30, color=C_WINE, bold=True,
            font_name='Microsoft YaHei')

# 左: 传统
x1, y1 = Inches(1.5), Inches(2.2)
shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x1, y1,
                                Inches(5), Inches(4.2))
shape.fill.solid()
shape.fill.fore_color.rgb = C_CREAM
shape.line.color.rgb = C_WINE
shape.line.width = Pt(1.5)

add_accent_bar(slide, Inches(1.8), Inches(2.5), Pt(4), Inches(1.0), C_WINE)
add_textbox(slide, Inches(2.2), Inches(2.5), Inches(4), Inches(0.5),
            "传统派 · 自然酒运动", font_size=22, color=C_WINE, bold=True,
            font_name='Microsoft YaHei')

txBox = add_textbox(slide, Inches(2.2), Inches(3.3), Inches(4), Inches(2.5),
                    "", font_size=15, color=C_DARK)
tf = txBox.text_frame
tf.word_wrap = True
bullets = ["最小干预酿造", "本地酵母自然发酵", "低二氧化硫使用",
           "强调风土表达", "反对技术标准化"]
for b in bullets:
    add_para(tf, f"·  {b}", font_size=16, color=C_DARK,
             font_name='Microsoft YaHei', line_spacing=1.6)

# 右: 现代
x2 = Inches(7)
shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x2, y1,
                                Inches(5), Inches(4.2))
shape.fill.solid()
shape.fill.fore_color.rgb = C_WINE
shape.line.fill.background()

add_accent_bar(slide, Inches(7.3), Inches(2.5), Pt(4), Inches(1.0), C_GOLD)
add_textbox(slide, Inches(7.7), Inches(2.5), Inches(4), Inches(0.5),
            "现代派 · 技术酿酒", font_size=22, color=C_GOLD, bold=True,
            font_name='Microsoft YaHei')

txBox = add_textbox(slide, Inches(7.7), Inches(3.3), Inches(4), Inches(2.5),
                    "", font_size=15, color=C_WHITE)
tf = txBox.text_frame
tf.word_wrap = True
bullets2 = ["筛选酵母精确控制", "温控不锈钢罐", "反渗透微氧技术",
            "稳定的国际化风格", "大规模标准化生产"]
for b in bullets2:
    add_para(tf, f"·  {b}", font_size=16, color=C_WHITE,
             font_name='Microsoft YaHei', line_spacing=1.6)

# 底部结论
add_textbox(slide, Inches(1.5), Inches(6.5), Inches(10), Inches(0.4),
            "两种方法之间的张力，本身就是一部人类技术哲学的变迁史",
            font_size=15, color=C_GRAY, alignment=PP_ALIGN.CENTER,
            font_name='Microsoft YaHei')

add_page_number(slide, 7)

# ════════════════════════════════════════════
# Slide 8: 情感载体
# ════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, C_BG_DARK)
add_shape_bg(slide, C_GOLD, top=0, height=Pt(4))

add_textbox(slide, Inches(1.5), Inches(0.5), Inches(4), Inches(0.5),
            "04 · EMOTIONAL VESSEL", font_size=12, color=C_GOLD, bold=True,
            font_name='Arial')
add_textbox(slide, Inches(1.5), Inches(0.9), Inches(10), Inches(0.7),
            "情感的载体：从个人记忆到集体无意识",
            font_size=30, color=C_GOLD, bold=True,
            font_name='Microsoft YaHei')

# 普鲁斯特金句框
shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
    Inches(1.5), Inches(2.0), Inches(10), Inches(2.0))
shape.fill.solid()
shape.fill.fore_color.rgb = C_WINE
shape.line.fill.background()

add_textbox(slide, Inches(2.0), Inches(2.2), Inches(9), Inches(1.5),
            '"带着点心渣的那一勺茶碰到我的上颚，\n顿时使我浑身一震……一种舒坦的快感传遍全身。"',
            font_size=20, color=C_GOLD, font_name='Microsoft YaHei',
            line_spacing=1.6)
add_textbox(slide, Inches(2.0), Inches(3.5), Inches(5), Inches(0.3),
            "— 马塞尔·普鲁斯特《追忆似水年华》",
            font_size=13, color=C_CREAM, font_name='Microsoft YaHei')

# 两列内容
y = Inches(4.5)
add_accent_bar(slide, Inches(1.5), y, Pt(4), Inches(2.0), C_GOLD)
txBox = add_textbox(slide, Inches(1.9), y, Inches(5), Inches(2.5),
                    "", font_size=16, color=C_CREAM)
tf = txBox.text_frame
tf.word_wrap = True
add_para(tf, "普鲁斯特效应", font_size=20, color=C_GOLD,
         bold=True, font_name='Microsoft YaHei', space_after=8)
add_para(tf, "嗅球与杏仁核、海马体之间存在直接神经连接，使气味能最直接地触发情感记忆。葡萄酒的数百种芳香化合物是普鲁斯特效应的绝佳媒介。",
         font_size=15, color=C_CREAM, line_spacing=1.5)

add_accent_bar(slide, Inches(7.5), y, Pt(4), Inches(2.0), C_GOLD)
txBox = add_textbox(slide, Inches(7.9), y, Inches(5), Inches(2.5),
                    "", font_size=16, color=C_CREAM)
tf = txBox.text_frame
tf.word_wrap = True
add_para(tf, "集体记忆之场", font_size=20, color=C_GOLD,
         bold=True, font_name='Microsoft YaHei', space_after=8)
add_para(tf, "香槟在庆祝活动中的符号化存在，使其超越饮料本身，成为胜利、喜悦与仪式的代名词——一种文化共识中的情感信号。",
         font_size=15, color=C_CREAM, line_spacing=1.5)

add_bottom_line(slide, C_GOLD)
add_page_number(slide, 8)

# ════════════════════════════════════════════
# Slide 9: 结论
# ════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, C_CREAM)
add_bottom_line(slide)

add_textbox(slide, Inches(1.5), Inches(0.5), Inches(4), Inches(0.5),
            "05 · CONCLUSION", font_size=12, color=C_GOLD, bold=True,
            font_name='Arial')
add_textbox(slide, Inches(1.5), Inches(0.9), Inches(10), Inches(0.7),
            "结论：一种新的阅读方式", font_size=30, color=C_WINE, bold=True,
            font_name='Microsoft YaHei')

# 三个核心维度
dims = [
    ("认识论", "挑战感官品鉴的霸权地位，将葡萄酒的理解从主观体验拓展为跨学科的知识实践"),
    ("方法论", "提供整合自然科学（气候学、地质学）与人文科学（历史学、文化研究）的分析路径"),
    ("实践", '赋予品酒行为新意义——每一次品鉴都是一次"阅读"，每一次阅读都是对过去的一次叩问'),
]

for i, (title, desc) in enumerate(dims):
    y = Inches(2.0) + i * Inches(1.5)
    # 编号圆
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.5), y + Pt(2),
                                    Inches(0.45), Inches(0.45))
    shape.fill.solid()
    shape.fill.fore_color.rgb = C_WINE
    shape.line.fill.background()
    tf = shape.text_frame
    tf.paragraphs[0].text = str(i + 1)
    tf.paragraphs[0].font.size = Pt(16)
    tf.paragraphs[0].font.color.rgb = C_WHITE
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.name = 'Arial'
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.word_wrap = False

    # 标题
    add_textbox(slide, Inches(2.3), y, Inches(2), Inches(0.4),
                title, font_size=20, color=C_WINE, bold=True,
                font_name='Microsoft YaHei')
    # 描述
    add_textbox(slide, Inches(2.3), y + Inches(0.45), Inches(9.5), Inches(0.8),
                desc, font_size=15, color=C_DARK, font_name='Microsoft YaHei',
                line_spacing=1.4)

# 底部金句
add_textbox(slide, Inches(1.5), Inches(6.2), Inches(10), Inches(0.7),
            "历史不仅写在书本里，躺在博物馆中，它还可以被种植、被酿造、被开启，最终被我们咽下。",
            font_size=18, color=C_WINE, bold=False, alignment=PP_ALIGN.CENTER,
            font_name='Microsoft YaHei', line_spacing=1.4)

add_page_number(slide, 9)

# ════════════════════════════════════════════
# Slide 10: 结束页
# ════════════════════════════════════════════
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, C_BG_DARK)
add_shape_bg(slide, C_GOLD, top=0, height=Pt(6))

add_textbox(slide, Inches(1.5), Inches(2.0), Inches(10), Inches(1.0),
            "感谢聆听", font_size=48, color=C_GOLD, bold=True,
            font_name='Microsoft YaHei', alignment=PP_ALIGN.CENTER)

add_textbox(slide, Inches(1.5), Inches(3.3), Inches(10), Inches(0.8),
            "LIQUID MEMORY", font_size=22, color=C_GOLD,
            font_name='Arial', alignment=PP_ALIGN.CENTER)

# 分隔线
shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,
    Inches(5.5), Inches(4.3), Inches(2.3), Pt(2))
shape.fill.solid()
shape.fill.fore_color.rgb = C_GOLD
shape.line.fill.background()

add_textbox(slide, Inches(1.5), Inches(4.7), Inches(10), Inches(0.6),
            "液体记忆：论葡萄酒作为历史、风土与情感的复合载体",
            font_size=16, color=C_CREAM, alignment=PP_ALIGN.CENTER,
            font_name='Microsoft YaHei')

add_textbox(slide, Inches(1.5), Inches(5.5), Inches(10), Inches(0.4),
            "完整论文阅读 · 论文.docx",
            font_size=13, color=C_GRAY, alignment=PP_ALIGN.CENTER,
            font_name='Microsoft YaHei')

add_textbox(slide, Inches(1.5), Inches(6.3), Inches(10), Inches(0.4),
            "2026",
            font_size=13, color=C_GRAY, alignment=PP_ALIGN.CENTER,
            font_name='Arial')

add_page_number(slide, 10)

# ── 保存 ──
output_path = r"C:\Users\Eric\Desktop\葡萄酒\液体记忆_葡萄酒论文PPT.pptx"
prs.save(output_path)
print(f"PPT 已生成: {output_path}")
print(f"共 {len(prs.slides)} 页")
