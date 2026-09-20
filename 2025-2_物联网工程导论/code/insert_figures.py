from pptx import Presentation
from pptx.util import Inches, Pt, Emu
import os

ppt_path = os.path.join(os.environ['USERPROFILE'], 'Desktop', '物联网工程', 'projects', 'RF_GPT_Presentation_ppt169_20260503', 'exports', 'RF_GPT_Presentation_20260504_144859.pptx')
output_path = os.path.join(os.environ['USERPROFILE'], 'Desktop', '物联网工程', 'RF-GPT_汇报幻灯片_含图_v2.pptx')
fig_dir = os.path.join(os.environ['USERPROFILE'], 'Desktop', '物联网工程', 'projects', 'RF_GPT_Presentation_ppt169_20260503', 'images', 'paper_figures')

prs = Presentation(ppt_path)

def add_img(slide, img_path, left_inch, top_inch, width_inch, height_inch=None):
    if not os.path.exists(img_path):
        print(f"  [SKIP] Image not found: {img_path}")
        return
    left = Inches(left_inch)
    top = Inches(top_inch)
    width = Inches(width_inch)
    if height_inch:
        slide.shapes.add_picture(img_path, left, top, width, Inches(height_inch))
    else:
        slide.shapes.add_picture(img_path, left, top, width)
    print(f"  [OK] {os.path.basename(img_path)}")

# Slide 5 (index 4) - Architecture: Fig 1 from page 4
add_img(prs.slides[4], os.path.join(fig_dir, 'page4-04.png'), 0.3, 1.6, 12.7, 4.2)

# Slide 7 (index 6) - Wireless Tech: Fig 2 from page 7
add_img(prs.slides[6], os.path.join(fig_dir, 'page7-07.png'), 0.3, 1.6, 12.7, 5.0)

# Slide 8 (index 7) - SKIP: already has paper figure embedded in SVG

# Slide 9 (index 8) - VLM vs RF-GPT: Fig 5 from page 11
add_img(prs.slides[8], os.path.join(fig_dir, 'page-11.png'), 0.3, 1.5, 12.7, 4.8)

# Slide 10 (index 9) - WBMC & WBOD: page 12
add_img(prs.slides[9], os.path.join(fig_dir, 'page-12.png'), 0.3, 1.5, 12.7, 4.8)

# Slide 11 (index 10) - WTR/WNUC/NRIE: page 13
add_img(prs.slides[10], os.path.join(fig_dir, 'page-13.png'), 0.3, 1.5, 12.7, 4.8)

# Slide 12 (index 11) - Ablation: page 14
add_img(prs.slides[11], os.path.join(fig_dir, 'page-14.png'), 0.3, 1.5, 12.7, 4.8)

prs.save(output_path)
print(f"\n[Done] Saved to: {output_path}")
